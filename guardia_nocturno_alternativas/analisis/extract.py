"""Lee Simulaciones.xlsx (valores calculados) y escribe datos.json para build_doc.js.

Uso: python3 extract.py [ruta/Simulaciones.xlsx] [ruta/datos.json]
"""
import json
import sys

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string as ci

XLSX = sys.argv[1] if len(sys.argv) > 1 else 'Simulaciones.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else 'datos.json'

wb = openpyxl.load_workbook(XLSX, data_only=True)
PRECIO = {'HT SICHERHEIT': 22800, 'SEPSISA': 24939.68, 'MACOTELA': 23000}
DIAS_PERIODO = 212  # 1-ene a 31-jul-2026
ROBOS_NOC_RED = 102  # robos en turno nocturno, ene-jul 2026
ROBOS_22_RED = 99  # robos de 22:00 a 22:59, ene-jul 2026


def hoja(nombre, cols):
    ws = wb[nombre]
    filas = []
    for r in range(2, 194):
        if ws.cell(r, 1).value is None:
            continue
        filas.append({c: ws.cell(r, ci(c)).value for c in cols})
    return pd.DataFrame(filas)


C1 = ['A', 'C', 'D', 'E', 'F', 'CK', 'CN', 'CU', 'CV', 'CW', 'CT', 'CY', 'CZ', 'DA', 'DD', 'DE', 'DF', 'DG']
C2 = ['A', 'CZ', 'DA', 'DD', 'DE', 'DF', 'DG', 'DH', 'EB', 'EC', 'ED', 'EE', 'EF', 'EG']
s1 = hoja('Sim 1 Conservador', C1)
s2 = hoja('Sim 2 Moderado', C2).add_suffix('_2')
d = pd.concat([s1, s2], axis=1)
assert (d.A == d.A_2).all() and len(d) == 191
d['p'] = d.F.map(PRECIO) * 12
for c in ['DD', 'DE', 'DF', 'DG', 'DD_2', 'DE_2', 'DF_2', 'DG_2', 'DH_2', 'EB_2', 'EC_2', 'ED_2', 'EE_2']:
    d[c] = d[c].fillna(0)
d['num'] = d.A.str[1:].astype(int)
d = d.sort_values('num')

# Base de captura (Semáforo)
ws = wb['Simulación']
sema = {}
for r in range(2, 194):
    a = ws.cell(r, 1).value
    if a:
        sema[a] = (sum((ws.cell(r, c).value or 0) for c in range(27, 31)), ws.cell(r, 31).value or 0)
d['intra'] = d.A.map(lambda a: sema[a][0])
d['arm'] = d.A.map(lambda a: sema[a][1])

base = {
    'farmacias': len(d), 'cdmx': int((d.CN == 'CDMX').sum()), 'foraneas': int((d.CN == 'Foránea').sum()),
    'intramuros': int(d.intra.sum()), 'armados': int(d.arm.sum()),
    'nocturnos': int(d.CY.sum()), 'costo_noc': float((d.CY * d.p).sum()),
    'empresas': [],
}
base['guardias'] = base['intramuros'] + base['armados']
for emp in ['HT SICHERHEIT', 'SEPSISA']:
    n = int(d[d.F == emp].intra.sum())
    base['empresas'].append({'empresa': emp, 'guardias': n, 'mensual': PRECIO[emp], 'anual': n * PRECIO[emp] * 12})
base['empresas'].append({'empresa': 'MACOTELA (armados)', 'guardias': base['armados'], 'mensual': PRECIO['MACOTELA'],
                         'anual': base['armados'] * PRECIO['MACOTELA'] * 12})
base['costo'] = sum(e['anual'] for e in base['empresas'])

# Farmacias que pierden el nocturno en cada alternativa
d['R1'] = d.DA == 'Sí'        # Conservador
d['R2'] = d.EG_2 == 'Sí'      # Moderado sin ventanilla anticipada
d['R3'] = d.EF_2 == 'Sí'      # Moderado con ventanilla anticipada
exp_cols = {1: ('DD', 'DE', 'DF', 'DG'), 2: ('DD_2', 'DE_2', 'DF_2', 'DG_2'), 3: ('DD_2', 'DE_2', 'DF_2', 'DG_2')}


def conservan(k):
    x = d[~d[f'R{k}'] & (d.CY > 0)]
    if k == 1:
        m = x.CZ
        return {
            'venta': int(m.str.contains('24 h').sum()),
            'alcaldia': int((~m.str.contains('24 h') & m.str.contains('Iztapalapa|GAM')).sum()),
            'bandas': int((m == 'Bandas nocturnas').sum()),
            'horario': int((m == 'Revisar horario').sum()),
            'total': len(x),
        }
    m = x.CZ_2
    cd, fo = x.CN == 'CDMX', x.CN == 'Foránea'
    return {
        'colonia': int((m == 'Colonia violenta (A)').sum()),
        'colonia_rev': int((m == 'A en revisión').sum()),
        'bandas_cdmx': int((cd & m.str.contains('Bandas nocturnas')).sum()),
        'bandas_for': int((fo & (m == 'Bandas nocturnas')).sum()),
        'venta_for': int((fo & m.str.contains('24 h')).sum()),
        'horario_cdmx': int((cd & (m == 'Revisar horario')).sum()),
        'horario_for': int((fo & (m == 'Revisar horario')).sum()),
        'robo22': int((m == 'Se quita con ventanilla anticipada').sum()),  # solo aplica en la alternativa 2
        'cdmx': int(cd.sum()), 'for': int(fo.sum()), 'total': len(x),
    }


alts = {}
for k in (1, 2, 3):
    x = d[d[f'R{k}']]
    dd, de, df, dg = exp_cols[k]
    ahorro = float(x.p.sum())
    alts[k] = {
        'farmacias': len(x), 'guardias': base['guardias'] - len(x), 'costo': base['costo'] - ahorro,
        'ahorro': ahorro, 'pct': ahorro / base['costo'],
        'cdmx': int((x.CN == 'CDMX').sum()), 'for': int((x.CN == 'Foránea').sum()),
        'ahorro_cdmx': float(x[x.CN == 'CDMX'].p.sum()), 'ahorro_for': float(x[x.CN == 'Foránea'].p.sum()),
        'ht': int((x.F == 'HT SICHERHEIT').sum()), 'sep': int((x.F == 'SEPSISA').sum()),
        'ahorro_ht': float(x[x.F == 'HT SICHERHEIT'].p.sum()), 'ahorro_sep': float(x[x.F == 'SEPSISA'].p.sum()),
        'noc_conservan': base['nocturnos'] - len(x),
        'rob_noc': int(x[dd].sum()), 'monto_noc': float(x[de].sum()),
        'rob_22': int(x[df].sum()), 'monto_22': float(x[dg].sum()),
        'rob_22_completo': int(x.DH_2.sum()),
        'ahorro_periodo': ahorro * DIAS_PERIODO / 365,
        'conservan': conservan(k),
    }
    alts[k]['veces'] = alts[k]['ahorro_periodo'] / alts[k]['monto_noc']
assert round(alts[1]['costo']) == 118013065 and round(alts[3]['costo']) == 109257865

red = {'rob_noc': ROBOS_NOC_RED, 'rob_22': ROBOS_22_RED, 'rob_22_completo': int(d.DH_2.sum())}


def fila(r, extra=None):
    o = {'clave': r.A, 'sucursal': r.C, 'alcaldia': r.D, 'formato': str(r.E), 'empresa': r.F}
    o.update(extra or {})
    return o


def motivo_alt1(m):
    partes = [{'24 h': 'Venta nocturna', 'Bandas nocturnas': 'robo organizado'}.get(p, p)
              for p in m.split(', ')]
    if len(partes) == 1:
        return partes[0][0].upper() + partes[0][1:]
    ult = partes[-1]
    conj = 'e' if ult.startswith('I') else 'y'
    return ', '.join(partes[:-1]) + f' {conj} ' + ult


comunes = [fila(r) for r in d[d.R1 & d.R3].itertuples()]
solo1 = [fila(r, {'rob_noc': int(r.DD), 'monto_noc': float(r.DE)}) for r in d[d.R1 & ~d.R3].itertuples()]
alt23 = []
for r in d[~d.R1 & d.R3].itertuples():
    alt23.append(fila(r, {
        'motivo1': motivo_alt1(r.CZ), 'rob22_completo': int(r.DH_2), 'rob22_2026': int(r.DF_2),
        'alternativas': '2 y 3' if r.R2 else 'Solo 3',
        'venta': r.CU == 'Sí', 'alcaldia_exc': isinstance(r.CV, str), 'bandas': (r.CW or 0) >= 2,
    }))
ventanilla = [x for x in alt23 if x['alternativas'] == 'Solo 3']


def grupo(mask, cols):
    x = d[mask]
    return {'n': len(x), 'rob_noc': int(x[cols[0]].sum()), 'monto_noc': float(x[cols[1]].sum()),
            'rob_22': int(x[cols[2]].sum()), 'rob_noc_completo': int(x.EB_2.sum()),
            'monto_noc_completo': float(x.EC_2.sum()), 'rob_22_completo': int(x.ED_2.sum()),
            'ahorro': float(x.p.sum())}


grupos = {
    'solo1': grupo(d.R1 & ~d.R3, ('DD', 'DE', 'DF')),
    'cambian37': grupo(~d.R1 & d.R3, ('DD_2', 'DE_2', 'DF_2')),
    'cambian34': grupo(~d.R1 & d.R2, ('DD_2', 'DE_2', 'DF_2')),
    'ventanilla3': grupo(d.R3 & ~d.R2, ('DD_2', 'DE_2', 'DF_2')),
}
x37 = d[~d.R1 & d.R3]
grupos['cambian37'].update({
    'venta': int((x37.CU == 'Sí').sum()),
    'alcaldia_sola': int(((x37.CU != 'Sí') & x37.CV.notna()).sum()),
    'bandas_sola': int(((x37.CU != 'Sí') & x37.CV.isna()).sum()),
})

# Sensibilidad: se recalcula con las columnas por farmacia (CDMX, robos organizados nocturnos y robos a las 22)
cd3 = d[(d.CN == 'CDMX') & d.R3]
G = 22800 * 12
sens = {
    'bandas2_alt3': alts[3]['ahorro'] - cd3[cd3.CW >= 2].p.sum(),
    'bandas2_alt2': alts[2]['ahorro'] - cd3[(cd3.CW >= 2) & cd3.R2].p.sum(),
    'robo22': {t: alts[3]['ahorro'] - cd3[(cd3.CU == 'Sí') & (cd3.DH_2 >= t)].p.sum() for t in (2, 3, 4, 5)},
}
assert round(sens['robo22'][3]) == round(alts[2]['ahorro'])

# Variante: alternativa 1 con el criterio de colonia como excepción
v = d[d.R1 & ~(d.CK == 'Sí')]
variante1 = {'guardias': base['guardias'] - len(v), 'ahorro': float(v.p.sum()), 'pct': float(v.p.sum()) / base['costo']}

# Estimación previa (Sim 1, filas 239 a 244)
ws1 = wb['Sim 1 Conservador']
previa = {'sin_bandas_horario': ws1['E239'].value, 'bandas': ws1['E240'].value, 'horario': ws1['E241'].value,
          'estimacion': ws1['E243'].value}

json.dump({'base': base, 'alts': alts, 'red': red, 'comunes': comunes, 'solo1': solo1, 'alt23': alt23,
           'ventanilla': ventanilla, 'grupos': grupos, 'sens': sens, 'variante1': variante1, 'previa': previa,
           'revision_horario': int((d.CT == 'Revisar').sum())},
          open(OUT, 'w'), ensure_ascii=False, indent=1, default=float)
print('ok', OUT)
